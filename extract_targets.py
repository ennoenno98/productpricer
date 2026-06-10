"""Extract country margin targets from the AP26 plan workbook.

Usage:
    python extract_targets.py path/to/AP26_Vanatari_official_Q1_Updated.xlsx

Reads the "Amazon Margins" tab:
    rows 3-8   headline per-country targets (COGS %, GP1/CM1, FBA %, GP2/CM2,
               demand capture %, channel margin/CM3)
    rows 76-88 monthly channel-margin (CM3) targets per country

Writes data/targets.json. The plan's "UK" is stored as "GB" to match the
Amazon fee report marketplace codes.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

DATA_DIR = Path(__file__).resolve().parent / "data"
COUNTRY_MAP = {"UK": "GB"}


def main(xlsx_path: str) -> None:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Amazon Margins"]

    # Headline block: country codes in row 2 (D2:J2), measures in C3:C8.
    countries = []
    for col in range(4, 27):
        v = ws.cell(row=2, column=col).value
        if v is None:
            break
        countries.append((col, COUNTRY_MAP.get(str(v).strip(), str(v).strip())))

    measures = {}
    for row in range(3, 9):
        name = str(ws.cell(row=row, column=3).value or "").strip().lower()
        measures[name] = row

    def grab(measure: str) -> dict[str, float]:
        row = measures[measure]
        return {c: round(float(ws.cell(row=row, column=col).value), 6) for col, c in countries}

    targets = {
        "source": Path(xlsx_path).name + " — 'Amazon Margins' tab",
        "cm1_target": grab("gp 1"),
        "cogs_target": grab("cogs"),
        "fba_pct_target": grab("fba"),
        "cm2_target": grab("gp2"),
        "demand_capture_target": grab("demand capture"),
        "cm3_target": grab("chanel margin (after demand capture)"),
    }

    # Monthly channel-margin block: header row 76 (C76..I76), months B77:B88.
    hdr_row = 76
    month_countries = []
    for col in range(3, 12):
        v = ws.cell(row=hdr_row, column=col).value
        if v is None:
            break
        month_countries.append((col, COUNTRY_MAP.get(str(v).strip(), str(v).strip())))
    monthly = {}
    for row in range(hdr_row + 1, hdr_row + 13):
        month = str(ws.cell(row=row, column=2).value or "").strip()
        if not month:
            continue
        monthly[month] = {
            c: round(float(ws.cell(row=row, column=col).value), 6)
            for col, c in month_countries
            if ws.cell(row=row, column=col).value is not None
        }
    targets["cm3_target_monthly"] = monthly

    DATA_DIR.mkdir(exist_ok=True)
    (DATA_DIR / "targets.json").write_text(json.dumps(targets, indent=2) + "\n")
    print(f"targets.json: {len(countries)} countries, {len(monthly)} months")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
