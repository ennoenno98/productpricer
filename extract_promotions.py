"""Extract promotion windows from the Seller Central promotions report.

Usage:
    python extract_promotions.py path/to/metricdata.xlsx

The report's ASIN column contains HYPERLINK formulas (=HYPERLINK(".../dp/<ASIN>";
"<ASIN>")), so ASINs are recovered from the formula text. Writes
data/promotions.csv with one row per promotion x ASIN: type, window,
promo price, original price. Marketplace attribution happens later in
estimate_elasticity.py by matching promo prices against daily implied prices.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd

DATA_DIR = Path(__file__).resolve().parent / "data"
ASIN_PAT = re.compile(r"/dp/([A-Z0-9]{10})")


def main(xlsx_path: str) -> None:
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    ws = wb.worksheets[0]
    header = [c.value for c in ws[1]]
    asin_idx = header.index("ASIN")

    asins = []
    for row in ws.iter_rows(min_row=2):
        v = row[asin_idx].value
        m = ASIN_PAT.search(str(v)) if v else None
        asins.append(m.group(1) if m else None)

    df = pd.read_excel(xlsx_path).iloc[: len(asins)].assign(asin=asins)
    df = df[(df["Promotion type"] != "Total") & df["asin"].notna()]
    out = (
        df[
            [
                "Promotion ID", "Promotion type", "asin", "Promotion price",
                "Original price", "Promotion start time", "Promotion end time",
            ]
        ]
        .rename(
            columns={
                "Promotion ID": "promotion_id",
                "Promotion type": "promotion_type",
                "Promotion price": "promo_price",
                "Original price": "original_price",
                "Promotion start time": "start",
                "Promotion end time": "end",
            }
        )
        .drop_duplicates()
    )
    out = out[out["start"].notna() & out["end"].notna()]
    DATA_DIR.mkdir(exist_ok=True)
    out.to_csv(DATA_DIR / "promotions.csv", index=False)
    print(
        f"promotions.csv: {len(out)} promo x ASIN rows, "
        f"{out['promotion_id'].nunique()} promotions, {out['asin'].nunique()} ASINs"
    )


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
